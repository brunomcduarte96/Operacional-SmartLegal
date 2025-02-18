from docx import Document
import os
from datetime import datetime
import locale
import pythoncom
import win32com.client

def gerar_procuracao(pasta_cliente, dados_cliente):
    """
    Gera a procuração preenchida com os dados do cliente
    """
    try:
        # Inicializar COM
        pythoncom.CoInitialize()
        
        # Configurar locale para português
        try:
            locale.setlocale(locale.LC_TIME, 'pt_BR.UTF-8')
        except locale.Error:
            try:
                locale.setlocale(locale.LC_TIME, 'Portuguese_Brazil.1252')
            except locale.Error:
                print("Não foi possível configurar o idioma para português. As datas serão exibidas no formato padrão.")
        
        # Obter caminhos absolutos
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.abspath(os.path.join(base_dir, "templates", "Modelo Procuracao JEC.docx"))
        pasta_cliente_abs = os.path.abspath(pasta_cliente)
        
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Arquivo modelo não encontrado em: {template_path}")
        
        # Carregar o modelo
        doc = Document(template_path)
        
        # Lista de placeholders e seus valores
        placeholders = {
            "{{nome}}": dados_cliente['nome'],
            "{{nacionalidade}}": dados_cliente['nacionalidade'],
            "{{estado_civil}}": dados_cliente['estado_civil'],
            "{{profissao}}": dados_cliente['profissao'],
            "{{rg}}": dados_cliente['rg'],
            "{{cpf}}": dados_cliente['cpf'],
            "{{rua}}": dados_cliente['rua'],
            "{{bairro}}": dados_cliente['bairro'],
            "{{complemento}}": dados_cliente['complemento'],
            "{{cep}}": dados_cliente['cep'],
            "{{cidade}}": dados_cliente['cidade'],
            "{{estado}}": dados_cliente['estado'],
            "{{data}}": dados_cliente['data']
        }
        
        # Substituir placeholders no documento
        for paragrafo in doc.paragraphs:
            for placeholder, valor in placeholders.items():
                if placeholder in paragrafo.text:
                    paragrafo.text = paragrafo.text.replace(placeholder, str(valor))
        
        # Salvar documento Word temporário com caminho absoluto
        docx_path = os.path.abspath(os.path.join(pasta_cliente_abs, f"Procuracao JEC - {dados_cliente['nome']}.docx"))
        doc.save(docx_path)
        
        # Converter para PDF usando Word COM
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False  # Ocultar o Word durante o processo
        
        try:
            doc = word.Documents.Open(docx_path)
            pdf_path = os.path.abspath(os.path.join(pasta_cliente_abs, f"Procuracao JEC - {dados_cliente['nome']}.pdf"))
            doc.SaveAs(pdf_path, FileFormat=17)  # 17 = PDF
            doc.Close()
        except Exception as e:
            raise Exception(f"Erro ao converter para PDF: {str(e)}\nCaminho do arquivo: {docx_path}")
        finally:
            word.Quit()
        
        # Remover arquivo Word temporário
        try:
            if os.path.exists(docx_path):
                os.remove(docx_path)
        except Exception as e:
            print(f"Aviso: Não foi possível remover o arquivo temporário: {str(e)}")
        
        return pdf_path
        
    except Exception as e:
        raise Exception(f"Erro ao gerar procuração: {str(e)}")
    finally:
        # Finalizar COM
        pythoncom.CoUninitialize() 